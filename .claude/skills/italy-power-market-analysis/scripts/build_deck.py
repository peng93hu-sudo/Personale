#!/usr/bin/env python3
"""Build the WiseGlow PUN deck by cloning the house template and replacing data.

Template-clone approach (NOT rebuild): the dark house style is per-shape, so we
open data/template_wiseglow.pptx and only swap chart data + dated text. This
guarantees the exact look. See references/deck-structure.md for the slide map.

Usage:
  python3 build_deck.py [--data data/market_data.json] [--date YYYY-MM-DD]
                        [--out output/WiseGlow_PUN_<date>.pptx]
"""
import argparse, datetime as dt, json, os
from pptx import Presentation
from pptx.chart.data import CategoryChartData

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(HERE, "data")
OUT = os.path.join(HERE, "output")


def replace_charts(prs, charts_by_key):
    replaced = 0
    for si, slide in enumerate(prs.slides, 1):
        ci = 0
        for shp in slide.shapes:
            if not shp.has_chart:
                continue
            ci += 1
            spec = charts_by_key.get((si, ci))
            if not spec:
                continue
            cd = CategoryChartData()
            cd.categories = spec["categories"]
            for s in spec["series"]:
                cd.add_series(s["name"], s["values"])
            shp.chart.replace_data(cd)
            replaced += 1
    return replaced


def apply_text(prs, subs):
    """subs: list of [old, new] literal substitutions across all text frames."""
    if not subs:
        return 0
    n = 0
    for slide in prs.slides:
        for shp in slide.shapes:
            if not shp.has_text_frame:
                continue
            for para in shp.text_frame.paragraphs:
                for run in para.runs:
                    for old, new in subs:
                        if old in run.text:
                            run.text = run.text.replace(old, new)
                            n += 1
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default=os.path.join(DATA, "market_data.json"))
    ap.add_argument("--template", default=os.path.join(DATA, "template_wiseglow.pptx"))
    ap.add_argument("--date", default=dt.date.today().isoformat())
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    payload = json.load(open(args.data, encoding="utf-8"))
    charts_by_key = {(c["slide"], c["chart_index"]): c for c in payload["charts"]}
    subs = payload.get("text_substitutions", [])

    prs = Presentation(args.template)
    n_charts = replace_charts(prs, charts_by_key)
    n_text = apply_text(prs, subs)

    os.makedirs(OUT, exist_ok=True)
    out = args.out or os.path.join(OUT, f"WiseGlow_PUN_{args.date}.pptx")
    prs.save(out)
    print(f"built {out}: {len(prs.slides)} slides, "
          f"{n_charts} charts replaced, {n_text} text runs updated")


if __name__ == "__main__":
    main()

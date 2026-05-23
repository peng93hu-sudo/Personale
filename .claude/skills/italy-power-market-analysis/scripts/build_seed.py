#!/usr/bin/env python3
"""Seed builder: vendor real Ember PUN + extract the template's chart data.

Run once to bootstrap the skill's data/ directory:
  - data/ember_italy_monthly_pun.csv   (real national monthly PUN, historical)
  - data/market_data.json              (every chart in the template, as data)

market_data.json is the contract build_deck.py consumes; fetchers overwrite
parts of it with fresh values before a build.
"""
import csv, json, os, sys
from pptx import Presentation
from pptx.oxml.ns import qn


def series_values(ser_el):
    """Read a series' numeric cache, preserving gaps (missing pt -> None)."""
    pts = ser_el.xpath(".//c:val//c:pt")
    cnt = ser_el.xpath(".//c:val//c:ptCount")
    n = int(cnt[0].get("val")) if cnt else len(pts)
    vals = [None] * n
    for pt in pts:
        idx = int(pt.get("idx"))
        v = pt.find(qn("c:v"))
        if v is not None and v.text not in (None, "") and 0 <= idx < n:
            try:
                vals[idx] = float(v.text)
            except ValueError:
                pass
    return vals

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(HERE, "data")
TEMPLATE = os.path.join(DATA, "template_wiseglow.pptx")


def vendor_ember(src):
    """Filter the Ember European CSV down to Italy monthly rows."""
    out = os.path.join(DATA, "ember_italy_monthly_pun.csv")
    rows = list(csv.reader(open(src, encoding="utf-8", errors="replace")))
    hdr = rows[0]
    ita = [r for r in rows[1:] if r[0].strip().lower() == "italy"]
    ita.sort(key=lambda r: r[2])
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(hdr)
        w.writerows(ita)
    print(f"ember -> {out} ({len(ita)} rows, {ita[0][2]}..{ita[-1][2]})")


def extract_charts(template):
    """Dump every chart's categories + series values, keyed by slide/chart idx."""
    prs = Presentation(template)
    charts = []
    for si, slide in enumerate(prs.slides, 1):
        ci = 0
        for shp in slide.shapes:
            if not shp.has_chart:
                continue
            ci += 1
            ch = shp.chart
            try:
                cats = [str(c) for c in ch.plots[0].categories]
            except Exception:
                cats = []
            series = []
            for s in ch.series:
                series.append({"name": s.name,
                               "values": series_values(s._element)})
            charts.append({
                "slide": si, "chart_index": ci,
                "type": str(ch.chart_type), "categories": cats,
                "series": series,
            })
    out = os.path.join(DATA, "market_data.json")
    payload = {
        "meta": {"source": "template seed (sample deck)",
                 "note": "real PUN F0 == Ember Italy monthly; refresh via fetchers"},
        "charts": charts,
    }
    json.dump(payload, open(out, "w", encoding="utf-8"),
              ensure_ascii=False, indent=2)
    print(f"charts -> {out} ({len(charts)} charts)")


if __name__ == "__main__":
    ember_src = sys.argv[1] if len(sys.argv) > 1 else None
    if ember_src and os.path.exists(ember_src):
        vendor_ember(ember_src)
    else:
        print("(skip ember: no source path given)")
    extract_charts(TEMPLATE)
